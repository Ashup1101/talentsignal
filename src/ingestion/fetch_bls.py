"""Pull BLS OEWS national occupation estimates into the S3 raw zone."""

from __future__ import annotations

import io
import logging
import sys
import zipfile
from datetime import date, datetime, timezone
from typing import Any

import requests
from openpyxl import load_workbook

from src.ingestion import s3_utils

logger = logging.getLogger(__name__)

# National OEWS release for May of year 20YY. Published each spring for the previous May.
OEWS_URL = "https://www.bls.gov/oes/special-requests/oesm{yy:02d}nat.zip"
REQUEST_TIMEOUT_S = 60
_REQUIRED_COLUMNS = ("OCC_CODE", "OCC_TITLE", "O_GROUP", "TOT_EMP", "A_MEDIAN")
# Footnote symbols BLS puts in place of numbers (see the workbook's "Field Descriptions" sheet):
# "*" wage not available, "**" employment not available, "#" wage at or above the top-code cap.
_SUPPRESSED = {"*", "**", "#"}


class BLSError(RuntimeError):
    """Raised when the BLS download fails or the workbook is not in the expected format."""


def _download(year: int) -> bytes | None:
    url = OEWS_URL.format(yy=year % 100)
    headers = {"User-Agent": f"talentsignal/0.1 ({s3_utils.require_env('BLS_CONTACT_EMAIL')})"}
    try:
        # BLS answers an unpublished year with a redirect to its home page, so don't follow it.
        resp = requests.get(url, headers=headers, timeout=REQUEST_TIMEOUT_S, allow_redirects=False)
    except requests.RequestException as exc:
        raise BLSError(f"Download of {url} failed: {exc}") from exc

    if resp.status_code in (301, 302, 404):
        logger.info("%s is not published (HTTP %d)", url, resp.status_code)
        return None
    if resp.status_code == 403:
        raise BLSError(
            f"BLS refused {url} (HTTP 403). BLS blocks requests whose User-Agent lacks contact "
            "details; check BLS_CONTACT_EMAIL is a real address."
        )
    if not resp.ok:
        raise BLSError(f"BLS returned HTTP {resp.status_code} for {url}")
    return resp.content


def _to_int(value: Any, field: str, occ_code: str) -> int | None:
    if isinstance(value, (int, float)):
        return int(value)
    if isinstance(value, str) and value.strip() in _SUPPRESSED:
        return None
    raise BLSError(f"Unexpected {field} value {value!r} for occupation {occ_code}")


def _parse_archive(archive: bytes, year: int) -> list[dict]:
    try:
        with zipfile.ZipFile(io.BytesIO(archive)) as zf:
            names = [n for n in zf.namelist() if n.lower().endswith(".xlsx")]
            if len(names) != 1:
                raise BLSError(f"Expected one .xlsx in the OEWS {year} archive, found {names}")
            workbook = load_workbook(io.BytesIO(zf.read(names[0])), read_only=True, data_only=True)
    except zipfile.BadZipFile as exc:
        raise BLSError(f"OEWS {year} download is not a valid zip archive") from exc

    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = [str(c).strip().upper() if c is not None else "" for c in next(rows, ())]
        missing = [c for c in _REQUIRED_COLUMNS if c not in header]
        if missing:
            raise BLSError(f"OEWS sheet {sheet.title!r} is missing columns {missing}; found {header}")
        col = {c: header.index(c) for c in _REQUIRED_COLUMNS}

        records = []
        for row in rows:
            if str(row[col["O_GROUP"]]).strip().lower() != "detailed":
                continue
            occ_code = str(row[col["OCC_CODE"]]).strip()
            records.append(
                {
                    "occupation_code": occ_code,
                    "title": str(row[col["OCC_TITLE"]]).strip(),
                    "total_employment": _to_int(row[col["TOT_EMP"]], "TOT_EMP", occ_code),
                    "median_wage": _to_int(row[col["A_MEDIAN"]], "A_MEDIAN", occ_code),
                    "reference_year": year,
                }
            )
    finally:
        workbook.close()

    if not records:
        raise BLSError(f"No detailed occupations found in the OEWS {year} workbook")
    return records


def fetch_bls_outlook(reference_year: int | None = None) -> list[dict]:
    """Download BLS OEWS national estimates and return one record per detailed occupation.

    Args:
        reference_year: OEWS reference year (May estimates). Defaults to the most
            recent published year, trying last year and then the year before.

    Returns:
        Dicts with occupation_code, title, total_employment, median_wage (annual,
        USD) and reference_year. total_employment / median_wage are None where BLS
        suppresses or top-codes the estimate.

    Raises:
        ValueError: If BLS_CONTACT_EMAIL is unset.
        BLSError: If no release is available or the download is malformed.
    """
    this_year = datetime.now(timezone.utc).year
    years = [reference_year] if reference_year else [this_year - 1, this_year - 2]

    for year in years:
        archive = _download(year)
        if archive is None:
            continue
        records = _parse_archive(archive, year)
        no_wage = sum(r["median_wage"] is None for r in records)
        logger.info(
            "Parsed %d detailed occupations from OEWS May %d (%d without an annual median wage)",
            len(records),
            year,
            no_wage,
        )
        return records

    raise BLSError(f"No national OEWS release found for {years}")


def save_to_s3(data: list[dict], run_date: date | None = None) -> str:
    """Write BLS records to s3://{S3_RAW_BUCKET}/bls/raw/{date}.json, never overwriting.

    Args:
        data: Records from fetch_bls_outlook.
        run_date: Partition date; defaults to today in UTC.

    Returns:
        The object's s3:// URI, whether newly written or already present.

    Raises:
        ValueError: If data is empty, or S3_RAW_BUCKET or an AWS env var is unset.
        S3OperationError: If an S3 call fails.
    """
    if not data:
        raise ValueError("Refusing to save an empty BLS dataset")

    bucket = s3_utils.require_env("S3_RAW_BUCKET")
    run_date = run_date or datetime.now(timezone.utc).date()
    key = f"bls/raw/{run_date.isoformat()}.json"
    uri = f"s3://{bucket}/{key}"

    if s3_utils.key_exists(bucket, key):
        logger.info("%s already exists; leaving it unchanged", uri)
        return uri

    s3_utils.upload_json(
        bucket,
        key,
        {
            "source": "bls_oews_national",
            "fetched_at": datetime.now(timezone.utc).isoformat(),
            "record_count": len(data),
            "records": data,
        },
    )
    return uri


def main() -> int:
    """Fetch the latest national OEWS release and land it in S3.

    Returns:
        Process exit code: 0 on success.
    """
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
    uri = save_to_s3(fetch_bls_outlook())
    logger.info("BLS data at %s", uri)
    return 0


if __name__ == "__main__":
    sys.exit(main())
